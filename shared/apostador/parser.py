"""Parser e validação de jogos para o módulo Apostador.

Suporta arquivos TXT/CSV (números separados por ';' ou ',', um jogo por linha)
e XLSX (1ª coluna, mesmas regras do Apostador C# original).
"""

from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Dict, Any


class JogoInvalidoError(Exception):
    """Linha do arquivo não pôde ser interpretada como jogo."""


class FormatoNaoSuportadoError(Exception):
    """Extensão de arquivo sem suporte."""


@dataclass
class JogoApostador:
    """Jogo de loteria com trevos (ex.: Mais Milionária: 6 dezenas + 2 trevos)."""
    numeros: List[int]
    trevos: List[int] = field(default_factory=list)


def _split_jogo(linha: str) -> List[int]:
    import re
    partes = re.split(r"[;,]", linha.strip())
    numeros: List[int] = []
    for parte in partes:
        texto = parte.strip()
        if not texto:
            continue
        try:
            numeros.append(int(texto))
        except ValueError:
            raise JogoInvalidoError(f"Número inválido: '{texto}'")
    return sorted(numeros)


def _parse_linha_com_trevos(linha: str) -> "JogoApostador":
    """Interpreta uma linha com dezenas e trevos.

    Formato: "<dezenas>|<trevos>" — ex.: "1,2,3,4,5,6|3,5".
    Linhas sem '|' são tratadas como apenas dezenas (trevos vazios).
    """
    if "|" in linha:
        dezenas_txt, trevos_txt = linha.split("|", 1)
        return JogoApostador(
            numeros=_split_jogo(dezenas_txt),
            trevos=_split_jogo(trevos_txt),
        )
    return JogoApostador(numeros=_split_jogo(linha))


def parse_jogos_texto(texto: str, tem_trevos: bool = False) -> List[Any]:
    """Interpreta o conteúdo de um arquivo de texto/CSV.

    Cada linha não vazia representa um jogo. Números separados por ';' ou ','.
    Uma linha que não seja numérica e não pareça jogo (ex.: header) é ignorada.

    tem_trevos=True (loterias com trevos): retorna List[JogoApostador] com
    dezenas e trevos separados; caso contrário retorna List[List[int]].
    """
    jogos: List[Any] = []
    for linha in texto.splitlines():
        linha_limpa = linha.strip()
        if not linha_limpa:
            continue
        if not any(c.isdigit() for c in linha_limpa):
            continue
        if tem_trevos:
            jogo = _parse_linha_com_trevos(linha_limpa)
            if not jogo.numeros and not jogo.trevos:
                continue
            jogos.append(jogo)
        else:
            numeros = _split_jogo(linha_limpa)
            if not numeros:
                continue
            jogos.append(numeros)
    return jogos


def parse_jogos_xlsx(caminho: str, tem_trevos: bool = False) -> List[Any]:
    """Lê jogos de uma planilha .xlsx (1ª coluna, números separados por ';')."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise FormatoNaoSuportadoError(
            "openpyxl não instalado. Instale com: pip install openpyxl"
        ) from exc

    wb = load_workbook(caminho, read_only=True, data_only=True)
    jogos: List[Any] = []
    try:
        ws = wb.active
        for linha in ws.iter_rows(min_col=1, max_col=1, values_only=True):
            valor = linha[0]
            if valor is None:
                continue
            texto = str(valor).strip()
            if not texto:
                continue
            if not any(c.isdigit() for c in texto):
                continue
            if tem_trevos:
                jogo = _parse_linha_com_trevos(texto)
                if jogo.numeros or jogo.trevos:
                    jogos.append(jogo)
            else:
                numeros = _split_jogo(texto)
                if numeros:
                    jogos.append(numeros)
    finally:
        wb.close()
    return jogos


def parse_arquivo_jogos(caminho: str, tem_trevos: bool = False) -> List[Any]:
    """Interpreta um arquivo de jogos pela extensão (.txt/.csv/.xlsx)."""
    ext = Path(caminho).suffix.lower()
    if ext in (".txt", ".csv"):
        texto = Path(caminho).read_text(encoding="utf-8", errors="replace")
        return parse_jogos_texto(texto, tem_trevos=tem_trevos)
    if ext == ".xlsx":
        return parse_jogos_xlsx(caminho, tem_trevos=tem_trevos)
    raise FormatoNaoSuportadoError(f"Extensão não suportada: '{ext}'")


def validar_jogos(
    cfg: Any,
    jogos: List[Any],
) -> Dict[str, Any]:
    """Valida uma lista de jogos contra a config da loteria.

    cfg deve expor: numero_minimo, numero_maximo, numeros_por_jogo.
    Para loterias com trevos (cfg.tem_trevos), cada jogo pode ser um
    JogoApostador (numeros + trevos) ou uma lista plana onde os últimos
    cfg.trevos_por_jogo valores são os trevos.

    Retorna:
        {"validos": [{"id", "numeros", "trevos"}], "invalids": [...]}
    """
    tem_trevos = getattr(cfg, "tem_trevos", False)
    trevos_por_jogo = getattr(cfg, "trevos_por_jogo", 0) or 0
    validos: List[Dict[str, Any]] = []
    invalidos: List[Dict[str, Any]] = []

    for idx, jogo in enumerate(jogos, start=1):
        if tem_trevos:
            if isinstance(jogo, JogoApostador):
                numeros = sorted(jogo.numeros)
                trevos = sorted(jogo.trevos)
            else:
                lista = list(jogo)
                if trevos_por_jogo and len(lista) > trevos_por_jogo:
                    trevos = sorted(lista[-trevos_por_jogo:])
                    numeros = sorted(lista[:-trevos_por_jogo])
                else:
                    numeros = sorted(lista)
                    trevos = []
        else:
            numeros = sorted(jogo)
            trevos = []
        erro = _validar_jogo(cfg, numeros)
        if erro is None and tem_trevos:
            erro = _validar_trevos(cfg, trevos)
        item: Dict[str, Any] = {"id": idx, "numeros": numeros, "trevos": trevos}
        if erro:
            item["mensagem"] = erro
            invalidos.append(item)
        else:
            validos.append(item)

    return {"validos": validos, "invalids": invalidos}


def _validar_trevos(cfg: Any, trevos: List[int]) -> str | None:
    qtd = len(trevos)
    if qtd != cfg.trevos_por_jogo:
        return f"Esperado {cfg.trevos_por_jogo} trevos, recebido {qtd}"
    if len(set(trevos)) != len(trevos):
        return "Trevo duplicado no jogo"
    fora = [t for t in trevos if not cfg.validar_trevo(t)]
    if fora:
        return f"Trevo fora da faixa ({cfg.trevos_min}-{cfg.trevo_maximo}): {fora[0]}"
    return None


def _validar_jogo(cfg: Any, numeros: List[int]) -> str | None:
    min_dezenas = cfg.numeros_por_jogo
    max_dezenas = getattr(cfg, "max_dezenas", None) or cfg.numeros_por_jogo
    min_val = cfg.numero_minimo
    max_val = cfg.numero_maximo

    qtd = len(numeros)
    if qtd < min_dezenas or qtd > max_dezenas:
        faixa = (
            f"{min_dezenas}" if min_dezenas == max_dezenas
            else f"{min_dezenas}-{max_dezenas}"
        )
        return f"Esperado {faixa} dezenas, recebido {qtd}"
    if len(set(numeros)) != len(numeros):
        return "Número duplicado no jogo"
    fora = [n for n in numeros if n < min_val or n > max_val]
    if fora:
        return f"Número fora da faixa ({min_val}-{max_val}): {fora[0]}"
    return None