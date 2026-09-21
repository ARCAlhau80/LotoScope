import pytest
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))


class TestParseJogosTxtCsv:
    def test_parse_linha_ponto_e_virgula(self):
        from shared.apostador.parser import parse_jogos_texto
        jogos = parse_jogos_texto("1;2;3;4;5;6;7;8;9;10;11;12;13;14;15")
        assert jogos == [[1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15]]

    def test_parse_linha_virgula(self):
        from shared.apostador.parser import parse_jogos_texto
        jogos = parse_jogos_texto("1,2,3,4,5,6,7,8,9,10,11,12,13,14,15")
        assert jogos == [[1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15]]

    def test_parse_multiplas_linhas(self):
        from shared.apostador.parser import parse_jogos_texto
        texto = "1;2;3;4;5;6;7;8;9;10;11;12;13;14;15\n1;2;3;4;5;6;7;8;9;10;11;12;13;14;16"
        jogos = parse_jogos_texto(texto)
        assert len(jogos) == 2

    def test_parse_linha_ordenada(self):
        from shared.apostador.parser import parse_jogos_texto
        jogos = parse_jogos_texto("15;1;7;3;9;2;10;4;11;5;12;6;13;8;14")
        assert jogos[0] == sorted([15, 1, 7, 3, 9, 2, 10, 4, 11, 5, 12, 6, 13, 8, 14])

    def test_parse_ignora_linhas_vazias(self):
        from shared.apostador.parser import parse_jogos_texto
        jogos = parse_jogos_texto("\n\n1;2;3\n\n4;5;6\n")
        assert len(jogos) == 2

    def test_parse_ignora_header(self):
        from shared.apostador.parser import parse_jogos_texto
        jogos = parse_jogos_texto("numeros\n1;2;3;4;5;6;7;8;9;10;11;12;13;14;15")
        assert jogos == [[1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15]]

    def test_parse_aceita_espacos(self):
        from shared.apostador.parser import parse_jogos_texto
        jogos = parse_jogos_texto(" 1 ; 2 ; 3 ; 4 ; 5 ; 6 ; 7 ; 8 ; 9 ; 10 ; 11 ; 12 ; 13 ; 14 ; 15 ")
        assert jogos == [[1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15]]

    def test_parse_linha_invalida_lanca_erro(self):
        from shared.apostador.parser import parse_jogos_texto, JogoInvalidoError
        with pytest.raises(JogoInvalidoError):
            parse_jogos_texto("1;2;abc;4")

    def test_parse_arquivo_vazio_retorna_lista_vazia(self):
        from shared.apostador.parser import parse_jogos_texto
        jogos = parse_jogos_texto("")
        assert jogos == []


class TestValidarJogos:
    def test_validar_jogo_valido_lotofacil(self):
        from shared.apostador.parser import validar_jogos
        from shared.lottery_config import LotteryConfig
        cfg = LotteryConfig(id="lotofacil", nome_jogo="Lotofácil", total_numeros=25, numeros_por_jogo=15)
        resultado = validar_jogos(cfg, [[1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15]])
        assert len(resultado["validos"]) == 1
        assert resultado["invalids"] == []
        assert resultado["validos"][0]["id"] == 1

    def test_validar_jogo_qtd_dezenas_errada(self):
        from shared.apostador.parser import validar_jogos
        from shared.lottery_config import LotteryConfig
        cfg = LotteryConfig(id="lotofacil", nome_jogo="Lotofácil", total_numeros=25, numeros_por_jogo=15)
        resultado = validar_jogos(cfg, [[1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14]])
        assert len(resultado["validos"]) == 0
        assert len(resultado["invalids"]) == 1
        assert "15" in resultado["invalids"][0]["mensagem"]

    def test_validar_jogo_numero_fora_faixa(self):
        from shared.apostador.parser import validar_jogos
        from shared.lottery_config import LotteryConfig
        cfg = LotteryConfig(id="lotofacil", nome_jogo="Lotofácil", total_numeros=25, numeros_por_jogo=15)
        resultado = validar_jogos(cfg, [[1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 30]])
        assert len(resultado["validos"]) == 0
        assert len(resultado["invalids"]) == 1
        assert "30" in resultado["invalids"][0]["mensagem"]

    def test_validar_jogo_numero_repetido(self):
        from shared.apostador.parser import validar_jogos
        from shared.lottery_config import LotteryConfig
        cfg = LotteryConfig(id="lotofacil", nome_jogo="Lotofácil", total_numeros=25, numeros_por_jogo=15)
        resultado = validar_jogos(cfg, [[1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 14]])
        assert len(resultado["validos"]) == 0
        assert len(resultado["invalids"]) == 1
        assert "duplicado" in resultado["invalids"][0]["mensagem"]

    def test_validar_mistura_validos_invalidos(self):
        from shared.apostador.parser import validar_jogos
        from shared.lottery_config import LotteryConfig
        cfg = LotteryConfig(id="lotofacil", nome_jogo="Lotofácil", total_numeros=25, numeros_por_jogo=15)
        resultado = validar_jogos(cfg, [
            [1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15],
            [1, 2, 3],
        ])
        assert len(resultado["validos"]) == 1
        assert len(resultado["invalids"]) == 1

    def test_jogo_16_dezenas_lotofacil_valido(self):
        from shared.apostador.parser import validar_jogos
        from shared.lottery_loader import get_loteria
        cfg = get_loteria("lotofacil")
        jogo_16 = list(range(1, 17))
        resultado = validar_jogos(cfg, [jogo_16])
        assert len(resultado["validos"]) == 1

    def test_lotofacil_aceita_ate_20_dezenas(self):
        from shared.apostador.parser import validar_jogos
        from shared.lottery_loader import get_loteria
        cfg = get_loteria("lotofacil")
        jogo_20 = list(range(1, 21))
        resultado = validar_jogos(cfg, [jogo_20])
        assert len(resultado["validos"]) == 1

    def test_lotofacil_rejeita_21_dezenas(self):
        from shared.apostador.parser import validar_jogos
        from shared.lottery_loader import get_loteria
        cfg = get_loteria("lotofacil")
        jogo_21 = list(range(1, 22))
        resultado = validar_jogos(cfg, [jogo_21])
        assert len(resultado["validos"]) == 0
        assert len(resultado["invalids"]) == 1
        assert "15-20" in resultado["invalids"][0]["mensagem"]


class TestParseXlsx:
    def test_parse_xlsx_com_openpyxl(self, tmp_path):
        from shared.apostador.parser import parse_jogos_xlsx
        try:
            import openpyxl
        except ImportError:
            pytest.skip("openpyxl não instalado")
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "1;2;3;4;5;6;7;8;9;10;11;12;13;14;15"
        ws["A2"] = "1;2;3;4;5;6;7;8;9;10;11;12;13;14;16"
        caminho = tmp_path / "jogos.xlsx"
        wb.save(caminho)
        jogos = parse_jogos_xlsx(str(caminho))
        assert len(jogos) == 2
        assert jogos[0] == [1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15]


class TestParseArquivo:
    def test_parse_txt(self, tmp_path):
        from shared.apostador.parser import parse_arquivo_jogos
        caminho = tmp_path / "jogos.txt"
        caminho.write_text("1;2;3;4;5;6;7;8;9;10;11;12;13;14;15", encoding="utf-8")
        jogos = parse_arquivo_jogos(str(caminho))
        assert len(jogos) == 1

    def test_parse_csv(self, tmp_path):
        from shared.apostador.parser import parse_arquivo_jogos
        caminho = tmp_path / "jogos.csv"
        caminho.write_text("1,2,3,4,5,6,7,8,9,10,11,12,13,14,15\n1,2,3,4,5,6,7,8,9,10,11,12,13,14,16", encoding="utf-8")
        jogos = parse_arquivo_jogos(str(caminho))
        assert len(jogos) == 2

    def test_parse_xlsx(self, tmp_path):
        from shared.apostador.parser import parse_arquivo_jogos
        try:
            import openpyxl
        except ImportError:
            pytest.skip("openpyxl não instalado")
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "1;2;3;4;5;6;7;8;9;10;11;12;13;14;15"
        caminho = tmp_path / "jogos.xlsx"
        wb.save(caminho)
        jogos = parse_arquivo_jogos(str(caminho))
        assert len(jogos) == 1

    def test_parse_extensao_desconhecida(self, tmp_path):
        from shared.apostador.parser import parse_arquivo_jogos, FormatoNaoSuportadoError
        caminho = tmp_path / "jogos.doc"
        caminho.write_text("1;2;3", encoding="utf-8")
        with pytest.raises(FormatoNaoSuportadoError):
            parse_arquivo_jogos(str(caminho))